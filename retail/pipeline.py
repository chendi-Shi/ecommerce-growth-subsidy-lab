"""Full-data, bounded-memory transaction analytics; run ``python -m retail``.

The financial ledger preserves anonymous customers and product returns. Customer
metrics use positive, identified purchases. These are descriptive retail findings,
not treatment effects or evidence about Douyin users.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import sys
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

import duckdb
from openpyxl import load_workbook
from retail.download import ranged_download

ROOT = Path(__file__).resolve().parents[1]
SOURCE_URL = "https://archive.ics.uci.edu/static/public/502/online+retail+ii.zip"
OFFICIAL_XLSX_URL = "https://archive.ics.uci.edu/ml/machine-learning-databases/00502/online_retail_II.xlsx"
MIRROR_COMMIT = "3a5eb9132055218af623a5550fc3f8d124b0e10c"
PINNED_MIRROR_URL = ("https://raw.githubusercontent.com/KristerMartinez/online_retail_ii/"
                     f"{MIRROR_COMMIT}/online_retail_II.xlsx")
# Fingerprint of the mirror bytes audited in this project, not an official UCI checksum.
AUDITED_WORKBOOK_SHA256 = "bcbe73b35f5b7babf197fb0cb983a11f5d9ff929078d4aa53d171b1f2df2e980"
SOURCE_PAGE = "https://archive.ics.uci.edu/dataset/502/online+retail+ii"
RAW_SCHEMA = """
invoice VARCHAR, stock_code VARCHAR, description VARCHAR, quantity DOUBLE,
invoice_date TIMESTAMP, unit_price DOUBLE, customer_id VARCHAR, country VARCHAR,
source_sheet VARCHAR, source_row BIGINT
"""
EXPORT_TABLES = (
    "monthly_growth", "monthly_decomposition", "monthly_cohort_retention",
    "repurchase_by_cohort", "repurchase_common_followup", "country_metrics", "price_band_metrics",
    "rfm_segment_validation",
)


def sql_path(path: Path) -> str:
    return "'" + path.resolve().as_posix().replace("'", "''") + "'"


def identifier(value: object) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        if not math.isfinite(float(value)):
            return None
        return str(int(value)) if float(value).is_integer() else str(value)
    return str(value).strip().upper() or None


def numeric(value: object) -> float | None:
    if value is None or value == "":
        return None
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (TypeError, ValueError):
        return None


def canonical_row(row: tuple, sheet: str, row_index: int) -> tuple:
    """Parse source types without imputing quantities, prices, IDs, or timestamps."""
    invoice, stock, description, quantity, timestamp, price, customer, country = row[:8]
    if isinstance(timestamp, date) and not isinstance(timestamp, datetime):
        timestamp = datetime.combine(timestamp, datetime.min.time())
    if not isinstance(timestamp, datetime):
        timestamp = None
    return (identifier(invoice), identifier(stock), description, numeric(quantity),
            timestamp.isoformat(sep=" ") if timestamp else None, numeric(price),
            identifier(customer), country, sheet, row_index)


def connect(database: Path | str = ":memory:") -> duckdb.DuckDBPyConnection:
    db = duckdb.connect(str(database))
    db.execute("SET memory_limit='256MB'")
    db.execute("SET threads=2")
    db.execute("SET preserve_insertion_order=false")
    return db


def acquire(raw_dir: Path) -> Path:
    """Acquire exactly the audited mirror version; UCI remains the source/citation."""
    raw_dir.mkdir(parents=True, exist_ok=True)
    workbook = raw_dir / "online_retail_II.xlsx"
    if workbook.exists() and zipfile.is_zipfile(workbook):
        if sha256(workbook) != AUDITED_WORKBOOK_SHA256:
            raise ValueError("Cached workbook differs from this project's audited input version; "
                             "use --xlsx with an explicitly sourced official workbook if intended")
        return workbook
    transport = ranged_download(PINNED_MIRROR_URL, workbook, parts=8)
    if transport["sha256"] != AUDITED_WORKBOOK_SHA256:
        raise ValueError("Downloaded mirror bytes differ from the version audited in this project")
    transport.update({"mirror_commit": MIRROR_COMMIT, "reproducible_download": PINNED_MIRROR_URL})
    write_json(workbook.with_suffix(".transport.json"), transport)
    return workbook


def sha256(path: Path) -> str:
    with path.open("rb") as source:
        return hashlib.file_digest(source, "sha256").hexdigest()


def load_rows(db: duckdb.DuckDBPyConnection, rows: Iterable[tuple], batch_path: Path,
              batch_size: int = 50_000) -> int:
    """Stream through bounded CSV batches; no million-row Python dataframe."""
    batch = []
    count = 0

    def flush() -> None:
        nonlocal count
        with batch_path.open("w", newline="", encoding="utf-8") as stream:
            csv.writer(stream, lineterminator="\n").writerows(batch)
        db.execute(f"COPY raw_lines FROM {sql_path(batch_path)} "
                   "(FORMAT CSV, HEADER false, NULL '', AUTO_DETECT false, "
                   "DELIMITER ',', QUOTE '\"', ESCAPE '\"', NEW_LINE '\\n')")
        count += len(batch)
        batch.clear()

    for row in rows:
        batch.append(row)
        if len(batch) >= batch_size:
            flush()
            print(f"Ingested {count:,} source invoice lines", flush=True)
    if batch:
        flush()
    batch_path.unlink(missing_ok=True)
    return count


def ingest_workbook(db: duckdb.DuckDBPyConnection, workbook: Path,
                    raw_dir: Path) -> dict:
    db.execute(f"CREATE OR REPLACE TABLE raw_lines ({RAW_SCHEMA})")
    sheets = {}
    # data_only avoids evaluating source formulas; official source has transaction values.
    with workbook.open("rb") as source:
        wb = load_workbook(source, read_only=True, data_only=True)
        for sheet in wb:
            stream = sheet.iter_rows(values_only=True)
            headers = next(stream)
            expected = ("Invoice", "StockCode", "Description", "Quantity",
                        "InvoiceDate", "Price", "Customer ID", "Country")
            if tuple(headers[:8]) != expected:
                raise ValueError(f"Unexpected source headers in {sheet.title}: {headers}")
            rows = (canonical_row(row, sheet.title, index)
                    for index, row in enumerate(stream, start=2) if any(v is not None for v in row))
            sheets[sheet.title] = load_rows(db, rows, raw_dir / "ingest_batch.csv")
        wb.close()
    return sheets


def build_marts(db: duckdb.DuckDBPyConnection) -> None:
    db.execute((ROOT / "sql" / "retail_marts.sql").read_text(encoding="utf-8"))


def one(db: duckdb.DuckDBPyConnection, query: str) -> dict:
    cursor = db.execute(query)
    names = [column[0] for column in cursor.description]
    return dict(zip(names, cursor.fetchone()))


def all_rows(db: duckdb.DuckDBPyConnection, query: str) -> list[dict]:
    cursor = db.execute(query)
    names = [column[0] for column in cursor.description]
    return [dict(zip(names, row)) for row in cursor.fetchall()]


def audit(db: duckdb.DuckDBPyConnection) -> dict:
    result = one(db, """
      SELECT count(*) AS source_rows,
        count(*) FILTER (WHERE customer_id IS NULL) AS source_missing_customer_rows,
        count(*) FILTER (WHERE invoice IS NULL OR stock_code IS NULL OR invoice_date IS NULL
             OR quantity IS NULL OR unit_price IS NULL) AS invalid_required_field_rows,
        count(*) FILTER (WHERE unit_price = 0) AS zero_price_rows,
        count(*) FILTER (WHERE unit_price < 0) AS negative_price_rows,
        count(*) FILTER (WHERE quantity = 0) AS zero_quantity_rows,
        count(*) FILTER (WHERE starts_with(invoice, 'C')) AS cancellation_rows,
        count(*) FILTER (WHERE quantity < 0) AS negative_quantity_rows,
        count(*) FILTER (WHERE starts_with(invoice, 'C') AND quantity > 0)
             AS cancellation_with_positive_quantity_rows,
        coalesce(sum(quantity * unit_price) FILTER
          (WHERE starts_with(invoice, 'C') AND quantity > 0 AND unit_price > 0),0)
             AS cancellation_positive_qty_raw_signed_amount_gbp,
        coalesce(sum(-2 * quantity * unit_price) FILTER
          (WHERE starts_with(invoice, 'C') AND quantity > 0 AND unit_price > 0),0)
             AS cancellation_positive_qty_sign_rule_delta_gbp,
        coalesce(sum(-2 * quantity * unit_price) FILTER
          (WHERE starts_with(invoice, 'C') AND quantity > 0 AND unit_price > 0
             AND regexp_full_match(stock_code, '[0-9]{5}[A-Z]*')),0)
             AS cancellation_positive_qty_product_ledger_delta_gbp
      FROM raw_lines
    """)
    result.update(one(db, """
      SELECT count(*) AS distinct_canonical_rows,
             sum(source_occurrences - 1) AS suspected_duplicate_rows,
             count(*) FILTER (WHERE source_occurrences > 1) AS duplicated_line_groups
      FROM canonical_lines
    """))
    result.update(one(db, """
      SELECT count(*) AS cross_sheet_duplicate_groups FROM (
        SELECT invoice, stock_code, description, quantity, invoice_date, unit_price,
               customer_id, country
        FROM raw_lines GROUP BY ALL HAVING count(DISTINCT source_sheet) > 1
      )
    """))
    result.update(one(db, """
      SELECT count(*) AS valid_clean_rows,
        count(*) FILTER (WHERE NOT is_product) AS non_product_rows,
        count(*) FILTER (WHERE is_product AND is_return) AS product_return_rows,
        count(*) FILTER (WHERE is_product AND NOT is_return) AS product_purchase_rows,
        count(*) FILTER (WHERE is_product AND customer_id IS NULL) AS anonymous_product_rows,
        sum(line_revenue) FILTER (WHERE is_product AND NOT is_return) AS gross_product_sales_gbp,
        -sum(line_revenue) FILTER (WHERE is_product AND is_return) AS product_returns_gbp,
        sum(line_revenue) FILTER (WHERE is_product) AS net_product_revenue_gbp,
        sum(line_revenue) FILTER (WHERE is_product AND customer_id IS NULL)
             AS anonymous_net_product_revenue_gbp,
        sum(line_revenue) AS all_codes_net_revenue_gbp,
        sum(line_revenue) FILTER (WHERE NOT is_product) AS non_product_net_revenue_gbp
      FROM retail_lines
    """))
    result.update(one(db, """
      SELECT count(*) AS purchase_orders, count(DISTINCT customer_id) AS observed_buyers,
        count(*) FILTER (WHERE customer_id IS NULL) AS anonymous_orders,
        sum(order_revenue) FILTER (WHERE customer_id IS NOT NULL) AS identified_purchase_revenue_gbp
      FROM retail_orders
    """))
    result["invoice_identity_audit"] = one(db, """
      WITH invoice_ids AS (
        SELECT invoice, count(DISTINCT customer_id) AS identified_ids,
          count(*) FILTER (WHERE customer_id IS NULL) AS anonymous_lines,
          count(DISTINCT country) AS countries
        FROM retail_lines WHERE is_product AND NOT is_return GROUP BY invoice
      )
      SELECT count(*) FILTER (WHERE identified_ids > 1) AS invoices_with_multiple_customer_ids,
        count(*) FILTER (WHERE identified_ids > 0 AND anonymous_lines > 0)
          AS invoices_mixing_anonymous_and_identified_lines,
        count(*) FILTER (WHERE countries > 1) AS invoices_with_multiple_countries
      FROM invoice_ids
    """)
    result["period"] = one(db, "SELECT * FROM observation")
    result["dedup_sensitivity"] = one(db, """
      SELECT sum(CASE WHEN starts_with(invoice, 'C') AND quantity > 0
                 THEN -quantity * unit_price ELSE quantity * unit_price END)
             AS product_net_revenue_without_dedup_gbp,
        sum(quantity * unit_price) FILTER
             (WHERE quantity > 0 AND NOT starts_with(invoice, 'C'))
             AS product_purchase_revenue_without_dedup_gbp
      FROM raw_lines WHERE invoice IS NOT NULL AND invoice_date IS NOT NULL
        AND regexp_full_match(stock_code, '[0-9]{5}[A-Z]*')
        AND quantity <> 0 AND unit_price > 0
    """)
    result["excluded_stock_codes"] = all_rows(db, """
      SELECT stock_code, min(description) AS representative_description,
        count(*) AS rows, sum(line_revenue) AS net_revenue_gbp,
        CASE WHEN stock_code IN ('D','M','POST','DOT','C2','BANK CHARGES','AMAZONFEE',
                                  'ADJUST','ADJUST2','CRUK')
             THEN 'charge_or_adjustment_code' ELSE 'unmapped_nonstandard_sku' END AS exclusion_class
      FROM retail_lines WHERE NOT is_product GROUP BY stock_code ORDER BY rows DESC
    """)
    result["unmapped_nonstandard_sku_sensitivity"] = one(db, """
      SELECT count(*) AS rows, coalesce(sum(line_revenue),0) AS net_revenue_gbp
      FROM retail_lines WHERE NOT is_product AND stock_code NOT IN
        ('D','M','POST','DOT','C2','BANK CHARGES','AMAZONFEE','ADJUST','ADJUST2','CRUK')
    """)
    result["customer_concentration"] = one(db, """
      WITH ranked AS (
        SELECT *, row_number() OVER (ORDER BY observed_purchase_revenue DESC) AS rank,
          count(*) OVER () AS buyers, sum(observed_purchase_revenue) OVER () AS total
        FROM customer_history
      )
      SELECT count(*) FILTER (WHERE rank <= ceil(buyers * .01)) AS top_one_percent_buyers,
        sum(observed_purchase_revenue) FILTER (WHERE rank <= ceil(buyers * .01))
          / max(total) AS top_one_percent_purchase_revenue_share
      FROM ranked
    """)
    return result


def verify(db: duckdb.DuckDBPyConnection) -> dict:
    checks = one(db, """
      SELECT max(abs(active_buyers - first_observed_buyers - repeat_buyers)) AS buyer_identity_error,
        max(abs(identified_purchase_revenue - active_buyers * orders_per_buyer * revenue_per_order))
          AS revenue_identity_error_gbp,
        max(abs(gross_product_sales - product_return_value - net_product_revenue))
          AS ledger_identity_error_gbp,
        max(abs(identified_purchase_revenue + coalesce(anonymous_purchase_revenue,0)
             - gross_product_sales)) AS customer_coverage_error_gbp
      FROM monthly_growth
    """)
    checks.update(one(db, """
      SELECT max(abs(reconciliation_error)) AS log_decomposition_error FROM monthly_decomposition
    """))
    checks.update(one(db, """
      SELECT count(*) AS invalid_retention_rows FROM monthly_cohort_retention
      WHERE retention_rate < 0 OR retention_rate > 1 OR active_customers > cohort_customers
    """))
    checks.update(one(db, """
      SELECT count(*) AS duplicate_order_keys FROM
      (SELECT customer_id, invoice FROM retail_orders GROUP BY 1,2 HAVING count(*)>1)
    """))
    for key, value in checks.items():
        if value is not None and abs(value) > (1e-6 if "error" in key else 0):
            raise AssertionError(f"Reconciliation failed: {key}={value}")
    return checks


def write_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str) + "\n",
                    encoding="utf-8")


def export(db: duckdb.DuckDBPyConnection, processed: Path, report_dir: Path) -> None:
    processed.mkdir(parents=True, exist_ok=True)
    report_dir.mkdir(parents=True, exist_ok=True)
    for table in EXPORT_TABLES:
        db.execute(f"COPY {table} TO {sql_path(report_dir / (table + '.csv'))} (HEADER, DELIMITER ',')")
    for table in ("retail_lines", "retail_orders", "customer_history", "rfm_holdout"):
        db.execute(f"COPY {table} TO {sql_path(processed / (table + '.parquet'))} "
                   "(FORMAT PARQUET, COMPRESSION ZSTD)")
    # Hive OpenCSVSerde does not support embedded newlines; the analysis parquet
    # preserves exact descriptions, only this transport sample replaces CR/LF.
    sample = """SELECT invoice, stock_code,
          regexp_replace(coalesce(description,''), '[\\r\\n]', ' ', 'g') AS description,
          quantity, invoice_date, unit_price, customer_id, country, source_sheet,
          line_revenue, is_product, is_return, observed_month
        FROM retail_lines ORDER BY invoice_date, invoice, stock_code, customer_id NULLS LAST,
          quantity, unit_price, description NULLS LAST, country, source_sheet LIMIT 50000"""
    db.execute(f"COPY ({sample}) TO {sql_path(processed / 'retail_hive_sample.csv')} (HEADER)")
    expected = one(db, f"""WITH sample AS ({sample}) SELECT count(*) AS rows,
      count(DISTINCT customer_id) AS identified_customers,
      sum(line_revenue) FILTER (WHERE is_product) AS net_product_revenue,
      sum(line_revenue) FILTER (WHERE is_product AND quantity > 0 AND unit_price > 0
             AND NOT is_return) AS product_purchase_amount FROM sample""")
    write_json(processed / "retail_hive_expected.json", expected)
    write_json(report_dir / "hive_sample_expected.json", expected)


def report(db: duckdb.DuckDBPyConnection, audit_result: dict, report_dir: Path) -> None:
    def money(value: float) -> str:
        return f"£{value:,.0f}"

    full = all_rows(db, "SELECT * FROM monthly_growth WHERE complete_month ORDER BY month")
    latest = full[-1]
    prior = next(row for row in full if row["month"] == "2010-11")
    mature = all_rows(db, """SELECT horizon_days, sum(eligible_customers) AS customers,
      sum(repurchased_customers) AS repurchases,
      sum(repurchased_customers)*1.0/sum(eligible_customers) AS rate
      FROM repurchase_by_cohort WHERE full_cohort_mature GROUP BY 1 ORDER BY 1""")
    segments = all_rows(db, "SELECT * FROM rfm_segment_validation")
    common = all_rows(db, "SELECT * FROM repurchase_common_followup")
    cohort_summary = all_rows(db, """SELECT month_index, sum(active_customers) AS active,
      sum(cohort_customers) AS eligible, sum(active_customers)*1.0/sum(cohort_customers) AS rate
      FROM monthly_cohort_retention WHERE month_index IN (1,3,6) GROUP BY 1 ORDER BY 1""")
    period = audit_result["period"]
    lines = [
        "# 真实交易分析：增长从哪里来、哪些复购机会值得实验",
        "",
        "本报告由 `python -m retail` 自动生成，源数据为英国线上礼品零售商交易，币种 GBP。"
        "原始数据属于真实交易记录；本报告不是抖音业务结论，也没有把交易观察关联解释成补贴因果效应。",
        "",
        "![Full-data retail overview](retail_overview.png)",
        "",
        f"来源：[UCI Online Retail II]({SOURCE_PAGE})，Chen, D. (2012)，DOI "
        "[10.24432/C5CG6D](https://doi.org/10.24432/C5CG6D)，CC BY 4.0。",
        "",
        f"实际处理 **{audit_result['source_rows']:,}** 条发票行，"
        f"{period['starts_on']} 至 {period['ends_on']}。最后一个月 2011-12 不完整，"
        "月度增长和月度留存均剔除该月；日级复购成熟度仍使用截至观测末日的记录。",
        "",
        "## 口径与审计",
        "",
        f"- 疑似完全重复行 {audit_result['suspected_duplicate_rows']:,} 条。"
        "按发票、SKU、描述、数量、日期、价格、客户、国家去重，不将来源工作表作为业务键。"
        "源文件无唯一发票行 ID，因此重复不必然是错误；保留去重前后金额敏感性。",
        f"- 原始客户 ID 缺失 {audit_result['source_missing_customer_rows']:,} 条；"
        "匿名交易参与商品销售/退货金额，排除买家数、RF 分群、复购分母。",
        f"- 清洗后 {audit_result['valid_clean_rows']:,} 条有效价格/数量行，"
        f"{audit_result['product_purchase_rows']:,} 条商品购买行、"
        f"{audit_result['product_return_rows']:,} 条商品退货/贷项行；"
        f"非商品代码 {audit_result['non_product_rows']:,} 条。",
        "- 商品 SKU 规则为五位数字加可选字母后缀。运费、折扣、银行费用、调整等独立代码"
        "不计入商品金额；不符合 SKU 规则的其余代码标为未映射，而不直接推断为费用。"
        "完整描述与代码金额敏感性见 `audit.json`；零价、负价、零数量及必需字段缺失另列。",
        "- 零价行可能是赠品或数据调整，因付费购买口径而排除，不直接认定为错误。"
        "这些审计标记可重叠，不能简单相加当成互斥清洗漏斗。",
        "- 购买订单为 `(customer_id, invoice)` 聚合的正向商品发票；退货不并入购买订单频次。"
        "取消标记或负数量归入退货账簿，无法判断实际退款原因或匹配原订单。",
        f"- {audit_result['cancellation_with_positive_quantity_rows']} 条 C 前缀且数量为正的记录，"
        f"原始 qty×price 合计 £{audit_result['cancellation_positive_qty_raw_signed_amount_gbp']:,.2f}；"
        "本项目按取消账簿符号规则转为负金额（口径假设），"
        f"相对原始符号的全代码净额变化为 £{audit_result['cancellation_positive_qty_sign_rule_delta_gbp']:,.2f}，"
        f"其中商品账簿变化为 £{audit_result['cancellation_positive_qty_product_ledger_delta_gbp']:,.2f}。",
        f"- 观察到 {audit_result['observed_buyers']:,} 位可识别买家、"
        f"{audit_result['purchase_orders']:,} 个购买发票（含匿名），"
        f"商品正向销售 {money(audit_result['gross_product_sales_gbp'])}、"
        f"退货金额 {money(audit_result['product_returns_gbp'])}、"
        f"净商品金额 {money(audit_result['net_product_revenue_gbp'])}。",
        f"- 不去重的净商品金额为 "
        f"{money(audit_result['dedup_sensitivity']['product_net_revenue_without_dedup_gbp'])}；"
        f"纳入所有有效非商品代码的净金额为 {money(audit_result['all_codes_net_revenue_gbp'])}。"
        "这些是收入口径，不是利润、增量 GMV 或 ROI。",
        "",
        "## 可识别用户增长拆解",
        "",
        "月购买金额 = 活跃买家 × 人均订单 × 单均金额。活跃买家 = 首次观察买家 + 连续月购买买家 + 回流买家。"
        "首次观察不等于平台新增注册用户；没有曝光、访问日志，不能计算完整获客/支付转化漏斗。",
        "",
        "| 同月比较 | 2010-11 | 2011-11 | 同比 |",
        "|---|---:|---:|---:|",
    ]
    for label, key, formatter in (
        ("活跃买家", "active_buyers", lambda x: f"{x:,.0f}"),
        ("首次观察买家", "first_observed_buyers", lambda x: f"{x:,.0f}"),
        ("复购买家", "repeat_buyers", lambda x: f"{x:,.0f}"),
        ("人均订单", "orders_per_buyer", lambda x: f"{x:.2f}"),
        ("单均金额", "revenue_per_order", money),
        ("可识别购买金额", "identified_purchase_revenue", money),
    ):
        lines.append(f"| {label} | {formatter(prior[key])} | {formatter(latest[key])} | "
                     f"{latest[key]/prior[key]-1:+.1%} |")
    lines += ["", "逐月对数变化精确拆解见 `monthly_decomposition.csv`。"
              "同比减少旺季季节性的干扰，但两年样本不足以识别因果增长驱动。",
              "", "## 复购机会及成熟度", "",
              "剔除观测初始三个月的首次观察 cohort，减轻左截断；仍无法确认用户此前从未购买。"
              "复购须发生在首次购买之后的另一个日历日，避免同日拆单。各窗口仅使用完整成熟 cohort：", "",
              "| 窗口 | 成熟首次观察买家 | 复购买家 | 累计复购率 |", "|---|---:|---:|---:|"]
    for row in mature:
        lines.append(f"| {row['horizon_days']} 天 | {row['customers']:,} | "
                     f"{row['repurchases']:,} | {row['rate']:.1%} |")
    lines += ["", "三个窗口成熟 cohort 集合不同，不能直接把表中差值解释为同一批用户的新增复购率。"
              "逐 cohort 结果及成熟标记见 `repurchase_by_cohort.csv`。", "",
              "固定使用已完整观测 90 天的 cohort，比较同一分母的累计复购："]
    for row in common:
        lines.append(f"- {row['horizon_days']} 天：{row['repurchase_rate']:.1%} "
                     f"（{row['repurchased_customers']:,}/{row['eligible_customers']:,}）。")
    lines += ["",
              "月留存定义为 cohort 在第 k 个日历月是否购买，允许跳月后回流，不是连续生存概率："]
    for row in cohort_summary:
        lines.append(f"- M{row['month_index']}：{row['rate']:.1%} "
                     f"（{row['active']:,}/{row['eligible']:,}，仅该偏移完整可观测 cohort）。")
    lines += ["", "## 人货场刻画与时间外验证", "",
              "人：2011-08-31 以前数据形成 RF 规则分群；2011-09 至 11 是未参与分群的未来评价窗口。"
              "频次与最近购买各分五档，按明确规则形成分群；历史金额未参与标签，用来刻画客户价值。"
              "相同值用 customer_id 稳定排序，"
              "分位边界存在人为划分，须做阈值敏感性；它是可解释基线，不是因果响应模型。", "",
              "| 分群 | 历史客户 | 历史人均购买金额 | 未来购买率 | 未来人均购买金额 |",
              "|---|---:|---:|---:|---:|"]
    for row in segments:
        lines.append(f"| {row['segment']} | {row['customers']:,} | "
                     f"{money(row['prior_revenue_per_customer'])} | {row['future_purchase_rate']:.1%} | "
                     f"{money(row['future_revenue_per_customer'])} |")
    concentration = audit_result["customer_concentration"]
    lines += ["", "货：按单品价格带比较购买人数、件数、金额，见 `price_band_metrics.csv`。"
              "买家和同一 SKU 可能跨价格带，人数/商品数不能跨组求和；没有商品类目，不编造品牌/品类。",
              "", "场：按国家市场划分，见 `country_metrics.csv`。该数据只有线上零售交易，"
              "没有短视频/直播/商城来源，因此国家是市场维度，不是流量场景或渠道归因。",
              "", f"已识别买家前 1%（{concentration['top_one_percent_buyers']:,} 人）贡献 "
              f"{concentration['top_one_percent_purchase_revenue_share']:.1%} 正向商品购买金额。"
              "来源说明中包含批发客户，金额分布受大户影响；业务建议需要普通买家/批发客分层实验。",
              "", "## 建议进入实验的业务假设", "",
              "1. 最近购买但低频客户存在二购培育空间。先做商品推荐/触达与无触达对照，主指标为"
              "预先定义 30 天复购率，次指标为人均净商品金额；对产品购买/退货共同观测。",
              "2. 高历史价值客户本来就有较高未来购买率，不能依据购买概率直接发补贴。"
              "先做客户级随机 holdout 识别净增量，避免把原本就会买的客户收入记作补贴收益。",
              "3. 沉睡客户按历史价值与最近购买分层，分别估计触达效果；该交易集没有触达、补贴"
              "和成本字段，不能用这里的 RF 分群表报告补贴 ROI。随机干预分析在独立真实 RCT 模块完成。",
              "", "## 复现与检查", "",
              "```powershell", "python -m retail", "python -m unittest tests.test_retail -v", "```", "",
              "默认下载使用固定 GitHub commit 的 UCI 数据镜像，以 8 个互不重叠字节块校验 "
              "ETag/ContentRange、完整文件长度、XLSX CRC 及项目已审计 SHA256。该 SHA256 是本项目输入指纹，"
              "不是 UCI 官方公布的校验和。已有匹配缓存及 `--reuse-db` 会复用本地数据。",
              "", "如手动从 UCI 获取官方文件，可用 `python -m retail --xlsx <path> "
              "--source-url <actual_download_url>`；流水线仍核对两工作表结构和总行数，记录实际指纹。", "",
              "原始 XLSX/ZIP、DuckDB 库和全量 Parquet 放在 gitignored `data/`；"
              "此目录中 CSV/audit 是全量数据派生输出。流水线使用只读 XLSX 行迭代、5 万行 CSV 批量入库、"
              "256 MB DuckDB 内存限制；不把百万行数据放入 pandas。",
              "", "`checks.json` 自动检查买家分类恒等式、收入乘法恒等式、商品销售/退货账簿、"
              "匿名/已识别金额衔接、逐月对数拆解及 cohort 分母约束。边界测试覆盖匿名客户、退货、"
              "同日拆单、右删失、去重跨工作表和价格无效值。",
              "", "Hive 使用确定性排序的 5 万行 CSV 运输样本；仅描述字段去除换行，分析 Parquet 不修改描述。"
              "真实 Hive 执行证据由仓库 CI 报告提供，不能把 DuckDB 执行说成 Hive 集群运行。", ""]
    (report_dir / "README.md").write_text("\n".join(lines), encoding="utf-8")


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-dir", type=Path, default=ROOT / "data" / "raw")
    parser.add_argument("--processed-dir", type=Path, default=ROOT / "data" / "processed")
    parser.add_argument("--report-dir", type=Path, default=ROOT / "reports" / "retail")
    parser.add_argument("--xlsx", type=Path, help="Already downloaded official workbook")
    parser.add_argument("--source-url", default=None,
                        help="Actual download URL to record when --xlsx was acquired separately")
    parser.add_argument("--reuse-db", action="store_true", help="Rebuild marts from existing ingestion")
    args = parser.parse_args(argv)
    args.raw_dir.mkdir(parents=True, exist_ok=True)
    args.processed_dir.mkdir(parents=True, exist_ok=True)
    args.report_dir.mkdir(parents=True, exist_ok=True)
    database = args.processed_dir / "retail.duckdb"
    if args.reuse_db and not (database.is_file() and
                             (args.processed_dir / "retail_source_manifest.json").is_file()):
        parser.error("--reuse-db requires a previously ingested database and source manifest")
    # Database-only refresh never contacts the network or needs the original XLSX.
    workbook = args.xlsx or (args.raw_dir / "online_retail_II.xlsx"
                            if args.reuse_db else acquire(args.raw_dir))
    db = connect(database)
    try:
        if not args.reuse_db:
            sheets = ingest_workbook(db, workbook, args.raw_dir)
            if sum(sheets.values()) != 1_067_371:
                raise ValueError(f"Official UCI source should have 1,067,371 rows; observed {sheets}")
            transport = (json.loads(workbook.with_suffix(".transport.json").read_text())
                         if workbook.with_suffix(".transport.json").exists() else None)
            write_json(args.processed_dir / "retail_source_manifest.json", {
                "source": SOURCE_PAGE,
                "download": args.source_url or (transport or {}).get("download", SOURCE_URL),
                "license": "CC BY 4.0",
                "workbook_sha256": sha256(workbook), "sheet_rows": sheets,
                "ingestion": "openpyxl read_only; 50000 row CSV batches; DuckDB 256MB limit",
                "transport": transport,
            })
        build_marts(db)
        audit_result = audit(db)
        checks = verify(db)
        manifest = json.loads((args.processed_dir / "retail_source_manifest.json").read_text())
        audit_result["source_manifest"] = manifest
        write_json(args.report_dir / "audit.json", audit_result)
        write_json(args.report_dir / "checks.json", checks)
        export(db, args.processed_dir, args.report_dir)
        report(db, audit_result, args.report_dir)
        print(json.dumps({"source_rows": audit_result["source_rows"],
                          "buyers": audit_result["observed_buyers"], "checks": checks}, indent=2))
    finally:
        db.close()
    # Release the database's memory before starting the plotting runtime.
    from retail.plots import render
    render(args.report_dir)


if __name__ == "__main__":
    main(sys.argv[1:])
